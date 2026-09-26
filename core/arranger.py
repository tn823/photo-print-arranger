"""
Core layout and document generation module for Photo Print Arranger.
Hỗ trợ cả chế độ in ảnh Chữ nhật truyền thống và ảnh Tròn dán vở Bé Ngoan.
"""
import os
import math
import shutil
import tempfile
from typing import List, Tuple, Optional, Callable, Dict
from PIL import Image

import docx
from docx.shared import Cm, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.enum.section import WD_ORIENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

from core.face_cropper import (
    CircleCropConfig,
    calculate_smart_circle,
    crop_circle_image
)

SUPPORTED_EXTENSIONS = ('.jpg', '.jpeg', '.png', '.bmp', '.webp')

def get_image_files(folder_or_files) -> List[str]:
    """Lấy danh sách các đường dẫn file ảnh hợp lệ."""
    if isinstance(folder_or_files, str):
        if os.path.isdir(folder_or_files):
            files = [
                os.path.join(folder_or_files, f)
                for f in os.listdir(folder_or_files)
                if f.lower().endswith(SUPPORTED_EXTENSIONS)
            ]
            return sorted(files)
        elif os.path.isfile(folder_or_files):
            return [folder_or_files] if folder_or_files.lower().endswith(SUPPORTED_EXTENSIONS) else []
        return []
    elif isinstance(folder_or_files, (list, tuple)):
        result = []
        for item in folder_or_files:
            result.extend(get_image_files(item))
        return sorted(list(dict.fromkeys(result)))
    return []

def get_auto_grid(images_per_page: int, orientation: str = 'landscape') -> Tuple[int, int]:
    """Tự động tính số hàng (rows) và số cột (cols) tối ưu cho số lượng ảnh chữ nhật/trang."""
    orientation = orientation.lower()
    mapping = {
        1: (1, 1),
        2: (1, 2) if orientation == 'landscape' else (2, 1),
        3: (1, 3) if orientation == 'landscape' else (3, 1),
        4: (2, 2),
        6: (2, 3) if orientation == 'landscape' else (3, 2),
        8: (2, 4) if orientation == 'landscape' else (4, 2),
        9: (3, 3),
        12: (3, 4) if orientation == 'landscape' else (4, 3),
        16: (4, 4),
    }
    if images_per_page in mapping:
        return mapping[images_per_page]
    
    cols = math.ceil(math.sqrt(images_per_page))
    rows = math.ceil(images_per_page / cols)
    if orientation == 'landscape' and rows > cols:
        rows, cols = cols, rows
    elif orientation == 'portrait' and cols > rows:
        rows, cols = cols, rows
    return rows, cols

def get_circle_grid(
    diameter_cm: float,
    orientation: str = 'portrait',
    margin_cm: float = 1.0,
    spacing_cm: float = 0.25
) -> Tuple[int, int]:
    """
    Tự động tính số hàng và cột tối đa cho ảnh tròn trên 1 trang A4.
    """
    orientation = orientation.lower()
    page_w, page_h = (29.7, 21.0) if orientation == 'landscape' else (21.0, 29.7)
    avail_w = page_w - (margin_cm * 2)
    avail_h = page_h - (margin_cm * 2) - 0.6  # Trừ lề an toàn chống tràn trang
    
    cols = max(1, int((avail_w + spacing_cm) / (diameter_cm + spacing_cm)))
    rows = max(1, int((avail_h + spacing_cm) / (diameter_cm + spacing_cm)))
    return rows, cols

class ImageArranger:
    def __init__(
        self,
        image_paths: List[str],
        output_docx: str,
        shape_mode: str = 'rectangle', # 'rectangle' hoặc 'circle'
        images_per_page: int = 2,
        circle_diameter_cm: float = 4.5,
        circle_configs: Optional[Dict[str, CircleCropConfig]] = None,
        fill_page: bool = False,
        repeat_count: int = 1,
        orientation: str = 'landscape', # 'landscape' hoặc 'portrait'
        margin_cm: float = 1.0,
        add_border: bool = False,
        export_pdf: bool = False,
        progress_callback: Optional[Callable[[int, int, str], None]] = None
    ):
        self.image_paths = [p for p in image_paths if os.path.exists(p)]
        self.output_docx = output_docx
        self.shape_mode = shape_mode.lower()
        self.images_per_page = max(1, images_per_page)
        self.circle_diameter_cm = max(1.5, float(circle_diameter_cm))
        self.circle_configs = circle_configs or {}
        self.fill_page = fill_page
        self.repeat_count = max(1, repeat_count)
        self.orientation = orientation.lower()
        self.margin_cm = margin_cm
        self.add_border = add_border
        self.export_pdf = export_pdf
        self.progress_callback = progress_callback

    def _notify(self, current: int, total: int, message: str):
        if self.progress_callback:
            self.progress_callback(current, total, message)

    def _apply_table_borders(self, table, force_none: bool = False):
        tblPr = table._tbl.tblPr
        if self.add_border and not force_none:
            # Đường cắt đứt mờ màu xám nhạt để dễ cắt
            tblBorders = parse_xml(
                '<w:tblBorders %s>'
                '<w:top w:val="dashed" w:sz="4" w:space="0" w:color="CCCCCC"/>'
                '<w:left w:val="dashed" w:sz="4" w:space="0" w:color="CCCCCC"/>'
                '<w:bottom w:val="dashed" w:sz="4" w:space="0" w:color="CCCCCC"/>'
                '<w:right w:val="dashed" w:sz="4" w:space="0" w:color="CCCCCC"/>'
                '<w:insideH w:val="dashed" w:sz="4" w:space="0" w:color="CCCCCC"/>'
                '<w:insideV w:val="dashed" w:sz="4" w:space="0" w:color="CCCCCC"/>'
                '</w:tblBorders>' % nsdecls('w')
            )
        else:
            tblBorders = parse_xml(
                '<w:tblBorders %s>'
                '<w:top w:val="none" w:sz="0" w:space="0" w:color="auto"/>'
                '<w:left w:val="none" w:sz="0" w:space="0" w:color="auto"/>'
                '<w:bottom w:val="none" w:sz="0" w:space="0" w:color="auto"/>'
                '<w:right w:val="none" w:sz="0" w:space="0" w:color="auto"/>'
                '<w:insideH w:val="none" w:sz="0" w:space="0" w:color="auto"/>'
                '<w:insideV w:val="none" w:sz="0" w:space="0" w:color="auto"/>'
                '</w:tblBorders>' % nsdecls('w')
            )
        tblPr.append(tblBorders)

    def _apply_cell_margins(self, cell, padding_pt: int = 4):
        tcPr = cell._tc.get_or_add_tcPr()
        pad_dxa = int(padding_pt * 20)
        tcMar = parse_xml(
            '<w:tcMar %s>'
            f'<w:top w:w="{pad_dxa}" w:type="dxa"/>'
            f'<w:left w:w="{pad_dxa}" w:type="dxa"/>'
            f'<w:bottom w:w="{pad_dxa}" w:type="dxa"/>'
            f'<w:right w:w="{pad_dxa}" w:type="dxa"/>'
            '</w:tcMar>' % nsdecls('w')
        )
        tcPr.append(tcMar)

    def calculate_optimal_size(self, rows: int, cols: int) -> Tuple[float, float]:
        """Tính kích thước tối ưu cho ảnh chữ nhật giữ trọn vẹn tỉ lệ."""
        if self.orientation == 'landscape':
            page_w, page_h = 29.7, 21.0
        else:
            page_w, page_h = 21.0, 29.7

        avail_w = page_w - (self.margin_cm * 2)
        avail_h = page_h - (self.margin_cm * 2)

        safety_h_margin = 0.6 if rows == 1 else (0.8 + (rows - 1) * 0.3)
        avail_h -= safety_h_margin

        cell_max_w = avail_w / cols
        cell_max_h = avail_h / rows

        aspect_ratios = []
        for p in self.image_paths[:10]:
            try:
                with Image.open(p) as img:
                    w, h = img.size
                    if h > 0:
                        aspect_ratios.append(w / h)
            except Exception:
                continue

        avg_ar = (sum(aspect_ratios) / len(aspect_ratios)) if aspect_ratios else 0.75

        target_h = cell_max_h * 0.98
        target_w = target_h * avg_ar

        if target_w > (cell_max_w * 0.98):
            target_w = cell_max_w * 0.98
            target_h = target_w / avg_ar

        return round(target_w, 2), round(target_h, 2)

    def generate(self) -> dict:
        total_images = len(self.image_paths)
        if total_images == 0:
            raise ValueError("Không tìm thấy ảnh hợp lệ nào để xử lý.")

        if self.shape_mode == 'circle':
            return self._generate_circle()
        else:
            return self._generate_rectangle()

    def _generate_rectangle(self) -> dict:
        total_images = len(self.image_paths)
        self._notify(0, total_images, "Đang khởi tạo tài liệu...")

        rows, cols = get_auto_grid(self.images_per_page, self.orientation)
        target_w, target_h = self.calculate_optimal_size(rows, cols)

        doc = docx.Document()
        for section in doc.sections:
            if self.orientation == 'landscape':
                section.orientation = WD_ORIENT.LANDSCAPE
                section.page_width = Cm(29.7)
                section.page_height = Cm(21.0)
            else:
                section.orientation = WD_ORIENT.PORTRAIT
                section.page_width = Cm(21.0)
                section.page_height = Cm(29.7)

            section.top_margin = Cm(self.margin_cm)
            section.bottom_margin = Cm(self.margin_cm)
            section.left_margin = Cm(self.margin_cm)
            section.right_margin = Cm(self.margin_cm)

        page_chunks = [
            self.image_paths[i:i + self.images_per_page]
            for i in range(0, total_images, self.images_per_page)
        ]
        total_pages = len(page_chunks)

        processed = 0
        for p_idx, chunk in enumerate(page_chunks):
            if p_idx > 0:
                doc.add_page_break()

            table = doc.add_table(rows=rows, cols=cols)
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            self._apply_table_borders(table)

            cell_width_cm = Cm(target_w + 0.3)

            for slot_idx in range(rows * cols):
                r = slot_idx // cols
                c = slot_idx % cols
                cell = table.cell(r, c)
                cell.width = cell_width_cm
                self._apply_cell_margins(cell, padding_pt=2)
                cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER

                p = cell.paragraphs[0]
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.0

                if slot_idx < len(chunk):
                    img_path = chunk[slot_idx]
                    run = p.add_run()
                    run.add_picture(img_path, height=Cm(target_h))
                    processed += 1
                    self._notify(processed, total_images, f"Đang thêm ảnh {processed}/{total_images}...")

        self._cleanup_trailing_paragraph(doc)

        os.makedirs(os.path.dirname(os.path.abspath(self.output_docx)), exist_ok=True)
        doc.save(self.output_docx)
        self._notify(total_images, total_images, f"Đã lưu Word: {os.path.basename(self.output_docx)}")

        pdf_path = None
        if self.export_pdf:
            self._notify(total_images, total_images, "Đang chuyển đổi sang PDF...")
            pdf_path = self.convert_to_pdf(self.output_docx)

        return {
            "total_images": total_images,
            "total_pages": total_pages,
            "grid": f"{rows}x{cols}",
            "image_size_cm": f"{target_w} x {target_h}",
            "docx_path": os.path.abspath(self.output_docx),
            "pdf_path": os.path.abspath(pdf_path) if pdf_path else None
        }

    def _generate_circle(self) -> dict:
        """Tạo trang in ảnh tròn dán sổ Bé Ngoan."""
        rows, cols = get_circle_grid(self.circle_diameter_cm, self.orientation, self.margin_cm)
        slots_per_page = rows * cols

        # Xác định danh sách ảnh cần in (in lặp lại hoặc in danh sách)
        if self.fill_page and len(self.image_paths) == 1:
            images_to_print = self.image_paths * slots_per_page
        elif self.repeat_count > 1:
            images_to_print = []
            for p in self.image_paths:
                images_to_print.extend([p] * self.repeat_count)
        else:
            images_to_print = list(self.image_paths)

        total_prints = len(images_to_print)
        self._notify(0, total_prints, "Đang chuẩn bị cắt ảnh tròn dán Bé Ngoan...")

        # Thư mục tạm để lưu ảnh tròn PNG sắc nét
        temp_dir = tempfile.mkdtemp(prefix="ppa_circle_")
        cropped_cache: Dict[str, str] = {}

        try:
            # Cắt ảnh tròn với AI và viền cắt kéo
            unique_paths = list(dict.fromkeys(images_to_print))
            for idx, p in enumerate(unique_paths):
                self._notify(idx, len(unique_paths), f"Đang căn chỉnh & cắt tròn ({idx + 1}/{len(unique_paths)})...")
                with Image.open(p) as raw_img:
                    # Lấy cấu hình crop đã lưu hoặc tự động tính toán
                    cfg = self.circle_configs.get(p)
                    if not cfg:
                        cfg = calculate_smart_circle(raw_img)

                    circle_pil = crop_circle_image(
                        raw_img, cfg,
                        output_size_px=800,
                        add_guide_border=self.add_border  # Sử dụng tùy chọn viền
                    )
                    cached_file = os.path.join(temp_dir, f"circ_{idx}.png")
                    circle_pil.save(cached_file, "PNG")
                    cropped_cache[p] = cached_file

            # Tạo tài liệu Word
            doc = docx.Document()
            page_w = 29.7 if self.orientation == 'landscape' else 21.0
            page_h = 21.0 if self.orientation == 'landscape' else 29.7

            for section in doc.sections:
                section.orientation = WD_ORIENT.LANDSCAPE if self.orientation == 'landscape' else WD_ORIENT.PORTRAIT
                section.page_width = Cm(page_w)
                section.page_height = Cm(page_h)
                section.top_margin = Cm(self.margin_cm)
                section.bottom_margin = Cm(self.margin_cm)
                section.left_margin = Cm(self.margin_cm)
                section.right_margin = Cm(self.margin_cm)

            avail_w = page_w - (self.margin_cm * 2)
            cell_w_cm = Cm(avail_w / cols)

            page_chunks = [
                images_to_print[i:i + slots_per_page]
                for i in range(0, total_prints, slots_per_page)
            ]
            total_pages = len(page_chunks)

            processed = 0
            for p_idx, chunk in enumerate(page_chunks):
                if p_idx > 0:
                    doc.add_page_break()

                table = doc.add_table(rows=rows, cols=cols)
                table.alignment = WD_TABLE_ALIGNMENT.CENTER
                # Không dùng viền bảng ô vuông để chỉ giữ viền tròn cắt kéo
                self._apply_table_borders(table, force_none=True)

                for slot_idx in range(slots_per_page):
                    r = slot_idx // cols
                    c = slot_idx % cols
                    cell = table.cell(r, c)
                    cell.width = cell_w_cm
                    self._apply_cell_margins(cell, padding_pt=1)
                    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER

                    p = cell.paragraphs[0]
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    p.paragraph_format.space_before = Pt(0)
                    p.paragraph_format.space_after = Pt(0)
                    p.paragraph_format.line_spacing = 1.0

                    if slot_idx < len(chunk):
                        orig_path = chunk[slot_idx]
                        circ_file = cropped_cache.get(orig_path)
                        if circ_file and os.path.exists(circ_file):
                            run = p.add_run()
                            run.add_picture(
                                circ_file,
                                width=Cm(self.circle_diameter_cm),
                                height=Cm(self.circle_diameter_cm)
                            )
                        processed += 1
                        self._notify(processed, total_prints, f"Đang thêm ảnh tròn {processed}/{total_prints}...")

            self._cleanup_trailing_paragraph(doc)

            os.makedirs(os.path.dirname(os.path.abspath(self.output_docx)), exist_ok=True)
            doc.save(self.output_docx)
            self._notify(total_prints, total_prints, f"Đã lưu Word: {os.path.basename(self.output_docx)}")

            pdf_path = None
            if self.export_pdf:
                self._notify(total_prints, total_prints, "Đang chuyển đổi sang PDF...")
                pdf_path = self.convert_to_pdf(self.output_docx)

            return {
                "total_images": total_prints,
                "total_pages": total_pages,
                "grid": f"{rows}x{cols} ({slots_per_page} ảnh/trang)",
                "image_size_cm": f"Ø {self.circle_diameter_cm} cm (Hình tròn)",
                "docx_path": os.path.abspath(self.output_docx),
                "pdf_path": os.path.abspath(pdf_path) if pdf_path else None
            }

        finally:
            shutil.rmtree(temp_dir, ignore_errors=True)

    def _cleanup_trailing_paragraph(self, doc):
        for p in doc.paragraphs:
            if not p.text and not p.runs:
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = Pt(1)
                run = p.add_run()
                run.font.size = Pt(1)

    @staticmethod
    def convert_to_pdf(docx_path: str, pdf_path: Optional[str] = None) -> str:
        """Chuyển đổi Word (.docx) sang PDF thông qua Microsoft Word Automation."""
        if not pdf_path:
            pdf_path = os.path.splitext(docx_path)[0] + ".pdf"

        try:
            import win32com.client
            word = win32com.client.Dispatch("Word.Application")
            word.Visible = False
            doc = word.Documents.Open(os.path.abspath(docx_path))
            doc.ExportAsFixedFormat(os.path.abspath(pdf_path), 17) # 17 = wdExportFormatPDF
            doc.Close(False)
            word.Quit()
            return pdf_path
        except Exception as e:
            raise RuntimeError(f"Lỗi khi xuất PDF qua Word: {e}")
